"""End-to-end P0 workflow on SQLite: sources → comparison → tasks → review → independent QC → field → exact export."""
import io
import zipfile
from datetime import date, datetime, timezone
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.auth import hash_password
from app.models.customer import CustomerDataset, CustomerPoi
from app.models.platform import ApprovedRecord, Task, User
from app.services.delivery import template_headers

LAT, LON = -6.1853, 106.8160  # inside Kampung Bali
TODAY = date.today().isoformat()
PASSWORD = "correct-horse-battery"


def make_user(Session, username, *roles):
    with Session() as db:
        db.add(User(username=username, display_name=username.title(), password_hash=hash_password(PASSWORD), roles=list(roles), aois=[]))
        db.commit()


def login(app, username):
    client = TestClient(app)
    assert client.post("/api/auth/login", json={"username": username, "password": PASSWORD}).status_code == 200
    return client


def here_poi(dataset_id, i, place_id, name, dlat=0.0, **raw):
    return CustomerPoi(dataset_id=dataset_id, place_id=place_id, name=name, category_code="600-6000-0061",
                       category_name="Convenience Store", country="IDN", address=f"Jalan Test {i}", latitude=LAT + dlat,
                       longitude=LON, routing_latitude=None, routing_longitude=None, display_point="POINT",
                       raw={"placeid": place_id, "name1": name, "primarycategory": "600-6000-0061", "status": "ACTIVE",
                            "admin3": "DKI Jakarta", "admin4": "Jakarta Pusat", "admin5": "Tanah Abang", "country": "IDN", **raw})


def csv_file(rows, header="source_record_id,name,category,latitude,longitude,phone,website,business_status,source_url,observed_on"):
    return ("\n".join([header] + rows)).encode()


@pytest.fixture
def world(sqlite_app):
    app, Session = sqlite_app
    for username, roles in [("ops", ["ops_manager"]), ("rev1", ["reviewer"]), ("rev2", ["reviewer", "qc_analyst"]),
                            ("qc", ["qc_analyst"]), ("field", ["field_verifier"])]:
        make_user(Session, username, *roles)
    dataset_id = uuid4()
    with Session() as db:
        db.add(CustomerDataset(id=dataset_id, filename="Kampung_Bali_IDN-6173.csv", customer="HERE", aoi="Kampung Bali",
                               country="IDN", sha256="a" * 64, row_count=3, columns=["placeid"]))
        db.flush()
        db.add_all([here_poi(dataset_id, 1, "here-1", "Toko Sejahtera", phone_numbers="0215550001"),
                    here_poi(dataset_id, 2, "here-2", "Warung Lama", dlat=0.0010),
                    here_poi(dataset_id, 3, "here-3", "Kedai Hilang", dlat=-0.0012)])
        db.commit()
    return app, Session, str(dataset_id)


def import_batch(client, name, source_type, rows):
    return client.post("/api/sources/batches", data={
        "aoi": "Kampung Bali", "source_type": source_type, "source_name": name,
        "license_reference": f"{name} licence 2026-10", "permitted_use": "Comparison and redistribution to HERE"},
        files={"file": (f"{name}.csv", csv_file(rows), "text/csv")})


def find(client, queue):
    tasks = client.get("/api/tasks/summary").json()["assignments"]
    return next((t for t in tasks if t["queue"] == queue), None)


def test_full_p0_workflow(world):
    app, Session, dataset_id = world
    anonymous = TestClient(app)
    assert anonymous.get("/api/tasks/summary").status_code == 401
    ops = login(app, "ops")

    # --- permitted sources; map-provider content is rejected ---
    rejected = import_batch(ops, "Scraped", "licensed_directory", [f"x1,Toko X,Shop,{LAT},{LON},,,,https://maps.google.com/?cid=1,"])
    assert rejected.status_code == 422 and "map provider" in str(rejected.json())
    assert import_batch(ops, "Brand sites", "official_website", [
        f"s1,Toko Sejahtera,Convenience Store,{LAT},{LON},0215550001,https://tokosejahtera.co.id,open,https://tokosejahtera.co.id/branches,{TODAY}",
        f"s2,Warung Lama,Convenience Store,{LAT + 0.0010},{LON},,,permanently closed,https://warunglama.id,{TODAY}",
        f"s3,Laundry Baru,Laundry,{LAT},{LON + 0.0005},,,open,https://laundrybaru.id,{TODAY}",
    ]).status_code == 201
    assert import_batch(ops, "Directory", "licensed_directory", [
        f"d1,TOKO SEJAHTERA,Convenience Store,{LAT + 0.00005},{LON},(021) 555-0001,,,https://directory.example/d1,{TODAY}",
    ]).status_code == 201

    # --- comparison: one master for the outlet seen in two sources; absence never deletes ---
    run = ops.post("/api/comparison/runs", json={"dataset_id": dataset_id}).json()
    assert run["stats"]["outcomes"] == {"no_change": 0, "potential_update": 1, "potential_new": 1, "verification_required": 1,
                                        "potential_delete": 1, "fallout_review": 0}
    assert run["stats"]["masters"] == 3 and run["stats"]["multi_source_masters"] == 1
    assert ops.post("/api/comparison/runs", json={"dataset_id": dataset_id}).json()["reused"] is True
    results = ops.get(f"/api/comparison/runs/{run['id']}/results").json()["results"]
    update = next(r for r in results if r["outcome"] == "potential_update")
    assert update["here_place_id"] == "here-1" and set(update["sources"]) == {"Brand sites", "Directory"}

    # --- idempotent task generation ---
    # Source-backed findings by default; verification of unobserved HERE POIs only when asked.
    assert ops.post(f"/api/comparison/runs/{run['id']}/tasks", json={}).json()["created"] == 3
    assert ops.post(f"/api/comparison/runs/{run['id']}/tasks", json={"include": ["verification_required"]}).json()["created"] == 1
    again = ops.post(f"/api/comparison/runs/{run['id']}/tasks", json={}).json()
    assert again["created"] == 0 and again["already_present"] == 3

    # --- allocation: no task reaches two reviewers ---
    rev1, rev2, qc, field = (login(app, u) for u in ("rev1", "rev2", "qc", "field"))
    first = {t["id"] for t in rev1.post("/api/tasks/claim", json={"limit": 2}).json()["claimed"]}
    second = {t["id"] for t in rev2.post("/api/tasks/claim", json={"limit": 10}).json()["claimed"]}
    assert len(first) == 2 and len(second) == 2 and not first & second
    assert rev1.post("/api/tasks/claim", json={"limit": 2}).json()["claimed"] == []  # kept in-progress, nothing left
    owner = {t["queue"]: (rev1 if t["id"] in first else rev2) for t in
             [*rev1.get("/api/tasks/summary").json()["assignments"], *rev2.get("/api/tasks/summary").json()["assignments"]]}

    # --- UPDATE: evidence is required, then independent QC approves ---
    reviewer = owner["review_update"]
    task = find(reviewer, "review_update")
    context = reviewer.get(f"/api/tasks/{task['id']}").json()
    assert context["here"]["place_id"] == "here-1" and len(context["observations"]) == 2
    assert any("Street View" in link["label"] for link in context["reference_links"])
    reviewer.post(f"/api/tasks/{task['id']}/start")
    payload = {"decision": "update", "verification_date": TODAY,
               "attributes": {"website": {"final": "https://tokosejahtera.co.id", "status": "corrected"}}}
    missing = reviewer.post(f"/api/tasks/{task['id']}/review", json=payload)
    assert missing.status_code == 422 and any("evidence" in i for i in missing.json()["detail"]["issues"])
    payload["evidence"] = [{"type": "official_website", "supports": "attribute", "fields": ["website"], "score": 0.9,
                            "url": "https://tokosejahtera.co.id/branches", "observed_on": TODAY}]
    assert reviewer.post(f"/api/tasks/{task['id']}/review", json=payload).status_code == 200

    # --- DELETE by rev2 (also a QC analyst): rev2 can never QC its own review ---
    delete_reviewer = owner["review_delete"]
    task = find(delete_reviewer, "review_delete")
    weak = {"decision": "delete", "delete_reason": "permanently_closed", "verification_date": TODAY,
            "evidence": [{"type": "official_website", "supports": "closure", "url": "https://warunglama.id/closed"}]}
    assert delete_reviewer.post(f"/api/tasks/{task['id']}/review", json=weak).status_code == 422  # E2 alone = 60 < 70
    weak["evidence"].append({"type": "official_social", "supports": "closure", "url": "https://facebook.com/warunglama/posts/1"})
    assert delete_reviewer.post(f"/api/tasks/{task['id']}/review", json=weak).status_code == 200

    rev2_qc = {t["queue"] for t in rev2.post("/api/tasks/claim", json={"limit": 10, "queues": ["qc_update", "qc_delete"]}).json()["claimed"]}
    if delete_reviewer is rev2:
        assert "qc_delete" not in rev2_qc
    for t in rev2.get("/api/tasks/summary").json()["assignments"]:
        if t["kind"] == "qc":
            rev2.post(f"/api/tasks/{t['id']}/release")
    claimed_qc = qc.post("/api/tasks/claim", json={"limit": 10}).json()["claimed"]
    for t in claimed_qc:
        if t["queue"] in ("qc_update", "qc_delete"):
            assert qc.post(f"/api/tasks/{t['id']}/qc", json={"action": "approve"}).status_code == 200

    # --- NEW: field review, back to independent QC, reclassify → rework → QC ---
    new_reviewer = owner["review_new"]
    task = find(new_reviewer, "review_new")
    assert new_reviewer.post(f"/api/tasks/{task['id']}/review", json={
        "decision": "field_review", "field_issue": "Confirm the laundry trades here and get the shop number"}).status_code == 200
    field_task = field.post("/api/tasks/claim", json={}).json()["claimed"][0]
    assert field_task["queue"] == "field"
    observed = datetime.now(timezone.utc).isoformat()
    assert field.post(f"/api/tasks/{field_task['id']}/field", json={
        "status": "existing", "observed_at": observed, "gps_latitude": LAT, "gps_longitude": LON + 0.0005, "gps_accuracy_m": 6,
        "signage_name": "Laundry Baru", "unit": "No. 12", "photos": ["https://evidence.jsan.example/f1.jpg"]}).status_code == 200
    qc_task = next(t for t in qc.post("/api/tasks/claim", json={}).json()["claimed"] if t["queue"] == "qc_fallout")
    assert qc.post(f"/api/tasks/{qc_task['id']}/qc", json={"action": "approve"}).status_code == 422
    assert qc.post(f"/api/tasks/{qc_task['id']}/qc", json={"action": "reclassify", "reclassified_to": "new",
                                                           "remarks": "Field confirms a new laundry"}).status_code == 200
    rework = find(new_reviewer, "rework")
    assert rework and "Reclassify to NEW" in rework["note"]
    new_payload = {"decision": "new", "verification_date": TODAY,
                   "attributes": {"name": {"final": "Laundry Baru", "status": "corrected"},
                                  "category_code": {"final": "100-1000-0000", "status": "corrected"},
                                  "latitude": {"final": LAT, "status": "corrected"},
                                  "longitude": {"final": LON + 0.0005, "status": "corrected"},
                                  "unit": {"final": "No. 12", "status": "corrected"}},
                   "evidence": [{"type": "field_observation", "supports": "existence", "url": "https://evidence.jsan.example/f1.jpg",
                                 "observed_on": TODAY, "score": 1},
                                {"type": "official_website", "supports": "existence", "url": "https://laundrybaru.id", "score": 0.8}]}
    assert new_reviewer.post(f"/api/tasks/{rework['id']}/review", json=new_payload).status_code == 200
    qc_new = next(t for t in qc.post("/api/tasks/claim", json={}).json()["claimed"] if t["queue"] == "qc_new")
    assert qc.post(f"/api/tasks/{qc_new['id']}/qc", json={"action": "approve"}).status_code == 200

    with Session() as db:
        approved = {a.action for a in db.query(ApprovedRecord)}
        assert approved == {"new", "update", "delete"}
        assert db.query(Task).filter(Task.kind == "qc", Task.status == "completed").count() == 4

    # --- delivery: gates, exact templates, customer files only in the package ---
    assert rev1.post(f"/api/delivery/datasets/{dataset_id}/snapshots").status_code == 403
    preflight = ops.get(f"/api/delivery/datasets/{dataset_id}/preflight").json()
    assert preflight["counts"] == {"new": 1, "update": 1, "delete": 1} and preflight["held_back"] == []
    snapshot = ops.post(f"/api/delivery/datasets/{dataset_id}/snapshots").json()
    files = {f["kind"]: f for f in snapshot["files"]}
    assert files["update"]["filename"].startswith("JSAN_ID_KampungBali_Update_") and files["update"]["filename"].endswith("_v1.0.xlsx")
    for action, expect in (("new", {"poi_name": "Laundry Baru", "category": "100-1000-0000"}),
                           ("update", {"uniqueID": "here-1", "WEBADDRESS": "https://tokosejahtera.co.id", "Final Action": "UPDATE"}),
                           ("delete", {"UniqueID": "here-2", "Final Action": "DELETE"})):
        wb = load_workbook(io.BytesIO(ops.get(f"/api/delivery/files/{files[action]['id']}").content))
        ws = wb.worksheets[0]
        title, headers = template_headers(action)
        assert ws.title == title
        assert [c.value for c in ws[1]][:len(headers)] == headers
        row = dict(zip(headers, [c.value for c in ws[2]]))
        assert all(row[k] == v for k, v in expect.items()), row
        assert row.get("Last Update") == TODAY
        assert ws.cell(row=3, column=1).value is None  # exactly one data row
    new_row = dict(zip(template_headers("new")[1], [c.value for c in load_workbook(io.BytesIO(
        ops.get(f"/api/delivery/files/{files['new']['id']}").content)).worksheets[0][2]]))
    assert new_row["unique_id"].startswith("JSAN-ID-")
    delete_reason = dict(zip(template_headers("delete")[1], [c.value for c in load_workbook(io.BytesIO(
        ops.get(f"/api/delivery/files/{files['delete']['id']}").content)).worksheets[0][2]]))["Decision/Delete Reason"]
    assert "Permanently closed" in delete_reason and "D6" in delete_reason
    package = zipfile.ZipFile(io.BytesIO(ops.get(f"/api/delivery/files/{files['package']['id']}").content))
    names = package.namelist()
    assert len(names) == 6 and not any("INTERNAL" in n for n in names)

    metrics = ops.get("/api/tasks/metrics/overview").json()
    assert metrics["qc_actions"]["approve"] == 3 and metrics["qc_rework_rate"] == 0.25


def test_reassignment_respects_roles_and_independence(world):
    app, Session, dataset_id = world
    ops = login(app, "ops")
    import_batch(ops, "Brand sites", "official_website", [f"s1,Toko Sejahtera,Shop,{LAT},{LON},0215550001,https://t.co.id,,,{TODAY}"])
    run = ops.post("/api/comparison/runs", json={"dataset_id": dataset_id}).json()
    ops.post(f"/api/comparison/runs/{run['id']}/tasks", json={})
    with Session() as db:
        users = {u.username: str(u.id) for u in db.query(User)}
        task_id = db.query(Task).filter(Task.queue == "review_update").one().id
    assert ops.post(f"/api/tasks/{task_id}/assign", json={"user_id": users["field"], "reason": "cover"}).status_code == 422
    assigned = ops.post(f"/api/tasks/{task_id}/assign", json={"user_id": users["rev2"], "reason": "workload"})
    assert assigned.status_code == 200 and assigned.json()["assignee"] == "Rev2"
    history = ops.get(f"/api/tasks/{task_id}").json()["history"]
    assert [h["event"] for h in history] == ["created", "reassigned"]
