"""HERE vs Google/Grab extract: New / Update / Delete / Matched / Review through the API, plus the Excel output."""
import io
import time
from uuid import uuid4

from openpyxl import Workbook, load_workbook

from app.models.customer import CustomerDataset
from app.services import area_collection
from tests.test_workflow_api import LAT, LON, here_poi, login, make_user


def extract_file(rows):
    """A provider extract workbook like the one 'Collect POIs (APIs)' downloads."""
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in ('Google', 'Grab'):
        ws = wb.create_sheet(sheet)
        ws.append(area_collection.HEADERS)
        for r in rows.get(sheet, []):
            ws.append([r.get(h, '') for h in area_collection.HEADERS])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def poi(placeid, name, dlat=0.0, **extra):
    return {'placeid': placeid, 'name1': name, 'displaylatitude': LAT + dlat, 'displaylongitude': LON, **extra}


def test_new_update_delete_matched_through_api(sqlite_app):
    app, Session = sqlite_app
    make_user(Session, 'ops', 'ops_manager')
    make_user(Session, 'rev', 'reviewer')
    dataset_id = uuid4()
    with Session() as db:
        db.add(CustomerDataset(id=dataset_id, filename='Kampung_Bali_IDN-6173.csv', customer='HERE', aoi='Kampung Bali',
                               country='IDN', sha256='b' * 64, row_count=4, columns=['placeid']))
        db.flush()
        db.add_all([here_poi(dataset_id, 1, 'here-1', 'Toko Sejahtera', phone_numbers='62215550001'),
                    here_poi(dataset_id, 2, 'here-2', 'Apotek Melati', dlat=0.0010),
                    here_poi(dataset_id, 3, 'here-3', 'Kedai Hilang', dlat=-0.0012),
                    here_poi(dataset_id, 4, 'here-4', 'Bengkel Jaya', dlat=0.0020)])
        db.commit()
    content = extract_file({
        'Google': [poi('g1', 'Toko Sejahtera', phone_numbers='62215550001'),                  # matched, no change
                   poi('g2', 'Apotek Melati', 0.0010, phone_numbers='62215559999'),            # update: phone
                   poi('g3', 'Kopi Baru Nusantara', 0.0030)],                                  # new
        'Grab': [poi('r1', 'Toko Sejahtera', 0.0001), poi('r4', 'Bengkel Jaya', 0.0020, status='CLOSED')],
    })
    assert login(app, 'rev').post('/api/provider-comparison/runs', data={'dataset_id': str(dataset_id)},
                                  files={'file': ('x.xlsx', content)}).status_code == 403
    client = login(app, 'ops')
    started = client.post('/api/provider-comparison/runs', data={'dataset_id': str(dataset_id)}, files={'file': ('x.xlsx', content)})
    assert started.status_code == 202, started.text
    run_id = started.json()['id']
    for _ in range(100):
        run = client.get(f'/api/provider-comparison/runs/{run_id}').json()
        if run['state'] != 'running':
            break
        time.sleep(0.05)
    assert run['state'] == 'finished', run
    assert run['stats']['counts'] == {'new': 1, 'update': 1, 'delete': 2, 'matched': 1, 'review': 0}
    assert run['stats']['delete_breakdown'] == {'Closed per provider': 1, 'Not found in Google or Grab': 1}

    def rows(kind):
        return client.get(f'/api/provider-comparison/runs/{run_id}/rows', params={'kind': kind}).json()['rows']
    assert rows('new')[0]['name1'] == 'Kopi Baru Nusantara'
    assert rows('update')[0]['here_placeid'] == 'here-2' and rows('update')[0]['changed_fields'] == 'phone'
    assert rows('matched')[0]['sources'] == 'Google Places, GrabMaps'  # Google and Grab copies merged into one row
    assert {r['here_placeid']: r['delete_basis'] for r in rows('delete')} == {'here-4': 'Closed per provider', 'here-3': 'Not found in Google or Grab'}
    detail = client.get(f'/api/provider-comparison/runs/{run_id}/rows/update/0').json()
    assert detail['proposed_phone'] == '62215559999'
    assert len(client.get(f'/api/provider-comparison/runs/{run_id}/map').json()['features']) == 5

    wb = load_workbook(io.BytesIO(client.get(f'/api/provider-comparison/runs/{run_id}/excel').content))
    assert wb.sheetnames == ['Summary', 'New', 'Update', 'Delete', 'Matched', 'Review']
    assert [c.value for c in wb['New'][1]][:35] == area_collection.HEADERS
    assert wb['New']['B2'].value == 'Kopi Baru Nusantara' and wb['Delete'].max_row == 3


def test_bad_upload_is_rejected(sqlite_app):
    app, Session = sqlite_app
    make_user(Session, 'ops', 'ops_manager')
    dataset_id = uuid4()
    with Session() as db:
        db.add(CustomerDataset(id=dataset_id, filename='d.csv', customer='HERE', aoi='Kampung Bali', country='IDN',
                               sha256='c' * 64, row_count=0, columns=[]))
        db.commit()
    r = login(app, 'ops').post('/api/provider-comparison/runs', data={'dataset_id': str(dataset_id)}, files={'file': ('x.xlsx', b'not excel')})
    assert r.status_code == 422


def test_extract_for_another_area_is_rejected(sqlite_app):
    app, Session = sqlite_app
    make_user(Session, 'ops', 'ops_manager')
    dataset_id = uuid4()
    with Session() as db:
        db.add(CustomerDataset(id=dataset_id, filename='d.csv', customer='HERE', aoi='Kampung Bali', country='IDN',
                               sha256='d' * 64, row_count=1, columns=[]))
        db.flush()
        db.add(here_poi(dataset_id, 1, 'here-1', 'Toko Sejahtera'))
        db.commit()
    manila = extract_file({'Google': [{'placeid': 'g', 'name1': 'Divisoria Mall', 'displaylatitude': 14.60, 'displaylongitude': 120.97}]})
    r = login(app, 'ops').post('/api/provider-comparison/runs', data={'dataset_id': str(dataset_id)}, files={'file': ('x.xlsx', manila)})
    assert r.status_code == 422 and 'does not cover Kampung Bali' in r.json()['detail']
