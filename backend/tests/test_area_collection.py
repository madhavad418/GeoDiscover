import io
import threading
import time
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.services import area_collection as service
from app.services.providers import ProviderError

POLYGON={'type':'Polygon','coordinates':[[[120,14],[120.01,14],[120.01,14.01],[120,14.01],[120,14]],
                                         [[120.004,14.004],[120.006,14.004],[120.006,14.006],[120.004,14.006],[120.004,14.004]]]}


def google(id,lon=120.002,lat=14.002):
    return {'id':id,'displayName':{'text':'=UNTRUSTED()'},'location':{'longitude':lon,'latitude':lat},
            'addressComponents':[{'types':['country'],'shortText':'PH'}],'attributions':[{'provider':'test'}]}


def job():
    return {'id':str(uuid4()),'name':'Test AOI','owner':'owner','created':time.time(),'state':'running',
            'queries':['restaurant'],'rows':{},'cancel':threading.Event(),'notice':'Not exhaustive',
            'providers':{p:{'state':'queued','requests':0,'records':0,'duplicates':0,'outside':0,
                           'invalid_results':0,'capped_queries':0,'remaining_queries':1,'error':None} for p in ('google','grab')}}


def test_polygon_holes_pagination_dedup_and_partial_failure(monkeypatch):
    calls=[]
    def search(provider,bounds,query,language,token=None):
        calls.append((provider,token))
        if provider=='grab': raise ProviderError('Provider returned HTTP 403.')
        if token: return [google('inside'),google('second')],None
        return [google('inside'),google('hole',120.005,14.005),google('outside',121,15)],'page2'
    monkeypatch.setattr(service,'search',search)
    j=job(); geom=service.validate_geometry(POLYGON)
    service.run(j,geom,[geom.bounds],10,'')
    assert j['state']=='finished'
    assert j['providers']['google']['records']==2
    assert j['providers']['google']['outside']==2
    assert j['providers']['google']['duplicates']==1
    assert ('google','page2') in calls
    assert j['providers']['grab']['state']=='failed'
    wb=load_workbook(io.BytesIO(service.workbook(j)))
    assert wb.sheetnames==['All POIs','Google','Grab']
    assert wb['All POIs'].max_row==wb['Google'].max_row+wb['Grab'].max_row-1
    assert [c.value for c in wb['Google'][1]]==service.HEADERS
    assert wb['Grab'].max_row==1
    assert '403' in wb['Grab']['A1'].comment.text
    assert wb['Google']['B2'].value=='=UNTRUSTED()'
    assert wb['Google']['B2'].data_type=='s'
    assert wb['Google']['I2'].value is None  # never invent routing coordinates


def test_budget_and_cancellation(monkeypatch):
    monkeypatch.setattr(service,'search',lambda *a: ([google('one')],'more'))
    j=job(); geom=service.validate_geometry(POLYGON)
    service.run(j,geom,[geom.bounds],1,'')
    assert j['providers']['google']['requests']==1
    assert j['providers']['google']['state']=='limited'
    j=job();j['cancel'].set()
    service.run(j,geom,[geom.bounds],10,'')
    assert j['state']=='cancelled'
    assert all(p['requests']==0 for p in j['providers'].values())


def test_invalid_and_oversized_geometry():
    for g in ({'type':'Point','coordinates':[120,14]},
              {'type':'Polygon','coordinates':[[[0,0],[10,0],[10,10],[0,0]]]}):
        with pytest.raises(ValueError):service.validate_geometry(g)


def test_owner_isolation():
    from app.api.area_collection import owned
    from fastapi import HTTPException
    from types import SimpleNamespace
    j=job();service.JOBS[j['id']]=j
    try:
        with pytest.raises(HTTPException) as err:owned(j['id'],SimpleNamespace(id='other'))
        assert err.value.status_code==404
        assert owned(j['id'],SimpleNamespace(id='owner')) is j
    finally:service.JOBS.pop(j['id'])


def test_grab_v2_storage_pagination_and_name(monkeypatch):
    calls=[]
    def request(bounds,query,language,token,storage):
        calls.append((token,storage))
        return {'ResultItems':[{'PlaceId':'grab-1','Title':'Local shop','Position':[120.002,14.002],
                               'Address':{'Label':'Shop address','Country':{'Code3':'PHL'}},
                               'Categories':[{'Name':'Shop','Primary':True}]}], 'NextToken':'next'}
    monkeypatch.setattr(service,'grab_search',request)
    items,token=service.search('grab',[120,14,120.01,14.01],'shop','fil','previous')
    assert calls==[('previous',True)]
    assert token=='next'
    row,_=service.normalized('grab',items[0])
    assert row['name1']=='Local shop'
    assert row['country']=='PHL'
    assert (row['primarycategory'],row['primarycategorynm'])==('600-6900-0000','Consumer Goods')  # HERE Find Places 'shop'
    assert row['suppliers']=='[GRABMAPS]' and row['chainid']=='0'


def test_here_formats_for_google_columns():
    item={'id':'g1','displayName':{'text':'Jollibee Divisoria'},'location':{'latitude':14.6012345678,'longitude':120.97},
          'businessStatus':'OPERATIONAL','primaryType':'fast_food_restaurant','types':['fast_food_restaurant','filipino_restaurant','restaurant'],
          'internationalPhoneNumber':'+63 2 8242 1109','addressComponents':[{'types':['country'],'shortText':'PH'},
          {'types':['locality'],'longText':'Manila'},{'types':['administrative_area_level_1'],'longText':'Metro Manila'}],
          'regularOpeningHours':{'periods':[{'open':{'day':d,'hour':8,'minute':30},'close':{'day':d,'hour':18,'minute':0}} for d in (1,2)]
                                 +[{'open':{'day':5,'hour':20,'minute':0},'close':{'day':6,'hour':2,'minute':0}}]}}
    row,note=service.normalized('google',item,'Restaurant')
    assert row['primarycategory'].startswith('100-') and row['primarycategorynm']
    assert (row['food_id'],row['food_name'])==('212-000','Filipino')
    assert (row['country'],row['admin2'],row['admin4'],row['status'])==('PHL','Metro Manila','Manila','ACTIVE')
    assert (row['phone_numbers'],row['mobile_numbers'])==('63282421109','')
    assert row['opening_hours']=='MONDAY:08:30:00-18:00:00,TUESDAY:08:30:00-18:00:00,FRIDAY:20:00:00-23:59:59,SATURDAY:00:00:00-02:00:00'
    assert row['chainname']=='Jollibee' and row['chainid'] not in ('','0')
    assert row['displaylatitude']==14.601235 and 'HERE category from provider category' in note
    assert service.phones(['+63 917 123 4567'])==('','639171234567')
    assert service.hours({'periods':[{'open':{'day':0,'hour':0,'minute':0}}]}).count('00:00:00-23:59:59')==7


def test_search_term_is_the_category_fallback():
    row,note=service.normalized('grab',{'PlaceId':'x','Title':'Somewhere','Position':[120,14],'Categories':[{'Id':'unknown_kind','Primary':True}]},'Pharmacy')
    assert row['primarycategorynm']=='Pharmacy' and 'search term' in note


def test_workbook_strips_control_characters():
    row={h:'' for h in service.HEADERS}; row.update(placeid='p',name1='18193 - Centerone',displaylatitude=1,displaylongitude=2)
    job={'id':'j','name':'A','created':0,'notice':'n','providers':{'google':{}},'rows':{'google':[(row,'note')]}}
    wb=load_workbook(io.BytesIO(service.workbook(job)))
    assert wb['Google']['B2'].value=='18193 - Centerone'
