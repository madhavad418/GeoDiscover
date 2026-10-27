"""Audit the supplied package without modifying source files.

Development dependencies: pypdf, openpyxl, shapely.
Usage: python audit_poc_material.py <POC Material directory>
"""
import argparse
import csv
import hashlib
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader
from shapely.geometry import Point, shape
from shapely.ops import unary_union


def audit(root):
    report = {'reviewed_on': '2026-10-27', 'files': [], 'aois': [], 'workbooks': [], 'documents': [], 'templates': []}
    for path in sorted(root.rglob('*')):
        if path.is_file():
            report['files'].append({'name': str(path.relative_to(root)).replace('\\', '/'), 'bytes': path.stat().st_size,
                                    'sha256': hashlib.sha256(path.read_bytes()).hexdigest()})
    category_codes = set()
    for path in sorted(root.glob('*.xlsx')):
        workbook = load_workbook(path, read_only=True, data_only=True)
        entry = {'file': path.name, 'sheets': []}
        for sheet in workbook:
            rows = sheet.iter_rows(values_only=True)
            header = list(next(rows, []))
            while header and header[-1] is None:
                header.pop()
            populated = 0
            errors = Counter()
            for row in rows:
                if not any(v is not None and str(v).strip() for v in row):
                    continue
                populated += 1
                for value in row:
                    if isinstance(value, str) and value.startswith(('#REF!', '#VALUE!', '#N/A', '#DIV/0!')):
                        errors[value] += 1
                if 'PDS Category Code' in header:
                    value = row[header.index('PDS Category Code')]
                    if value and str(value)[0].isdigit():
                        category_codes.add(str(value))
            entry['sheets'].append({'name': sheet.title, 'columns': header, 'populated_rows': populated,
                                    'excel_errors': dict(errors), 'visibility': sheet.sheet_state})
        workbook.close()
        report['workbooks'].append(entry)
        if path.name.startswith('Supplier Name'):
            report['templates'].append(entry)
    ids = Counter()
    for path in sorted((root / 'Requested_AOI_Places' / 'data_extract').glob('*.csv')):
        stem = 'Divisoria_PHL' if 'Divisoria' in path.name else 'Kampung_Bali_IDN' if 'Kampung' in path.name else 'Thanon_Phaya_Thai_THA'
        boundary_path = root / 'Requested_AOI_Places' / 'boundary' / f'{stem}.geojson'
        boundary = json.loads(boundary_path.read_text(encoding='utf-8-sig'))
        geom = unary_union([shape(f['geometry']) for f in boundary['features']])
        with path.open(encoding='utf-8-sig', newline='') as f:
            reader = csv.DictReader(f)
            rows = list(reader)
            columns = reader.fieldnames
        local_ids = Counter(r['placeid'] for r in rows)
        ids.update(local_ids)
        outside, invalid, missing_name, missing_category = 0, 0, 0, 0
        for row in rows:
            missing_name += not bool(row.get('name1', '').strip())
            missing_category += not bool(row.get('primarycategory', '').strip())
            try:
                lat, lon = float(row['displaylatitude']), float(row['displaylongitude'])
                if not (-90 <= lat <= 90 and -180 <= lon <= 180):
                    raise ValueError()
                outside += not geom.covers(Point(lon, lat))
            except (ValueError, TypeError):
                invalid += 1
        unknown = Counter(r['primarycategory'] for r in rows if r['primarycategory'] not in category_codes)
        report['aois'].append({'name': stem.rsplit('_', 1)[0].replace('_', ' '), 'country': stem[-3:],
            'file': path.name, 'boundary_file': boundary_path.name, 'rows': len(rows), 'columns': columns,
            'unique_ids': len(local_ids), 'duplicate_ids': sum(n - 1 for n in local_ids.values()),
            'missing_names': missing_name, 'missing_category': missing_category, 'invalid_coordinates': invalid,
            'outside_boundary': outside, 'boundary_valid': geom.is_valid, 'bbox': list(geom.bounds),
            'unknown_category_records': sum(unknown.values()), 'unknown_categories': dict(unknown),
            'category_count': len({r['primarycategory'] for r in rows}),
            'verification_date_field_present': any('verif' in c.lower() for c in columns)})
    report['cross_file_duplicate_ids'] = sum(n - 1 for n in ids.values())
    report['category_reference_codes'] = len(category_codes)
    report['category_codes'] = sorted(category_codes)
    for path in sorted(root.glob('*.pdf')):
        pdf = PdfReader(path)
        report['documents'].append({'file': path.name, 'pages': len(pdf.pages),
            'extracted_characters': sum(len(page.extract_text() or '') for page in pdf.pages)})
    for path in sorted(root.rglob('*.qmd')):
        # Include the metadata's CRS declaration as an inspected input, not a POI source.
        import xml.etree.ElementTree as ET
        xml = ET.fromstring(path.read_text(encoding='utf-8-sig'))
        report.setdefault('qgis_metadata', []).append({'file': path.name,
            'crs': [e.text for e in xml.iter('authid') if e.text]})
    return report


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('folder', type=Path)
    args = parser.parse_args()
    result = audit(args.folder)
    out = Path(__file__).parent / 'app' / 'resources' / 'poc_audit.json'
    out.parent.mkdir(exist_ok=True)
    out.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding='utf-8')
    print(json.dumps({k: result[k] for k in ('aois','cross_file_duplicate_ids','category_reference_codes','documents')}, indent=2))
    print('Report:', out)
