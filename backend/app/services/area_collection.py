"""Bounded provider searches for uploaded polygons; temporary, owner-scoped exports."""
import io
import json
import math
import re
import threading
import time
from collections import deque
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import quote, urlencode
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from shapely.geometry import Point, box, shape
from shapely.errors import ShapelyError

from ..config import get_settings
from .providers import ProviderError, request_json, grab_search
from .matching import norm_name
from .reference import categories, category_by_name, chain_ids, chain_names, find_places, foodtype_by_name, foodtypes

HEADERS = json.loads((Path(__file__).resolve().parents[1] / 'resources/poc_audit.json').read_text(encoding='utf-8'))['aois'][0]['columns']
FIELDS = 'nextPageToken,places.id,places.displayName,places.location,places.formattedAddress,places.addressComponents,places.primaryType,places.primaryTypeDisplayName,places.types,places.businessStatus,places.nationalPhoneNumber,places.internationalPhoneNumber,places.websiteUri,places.regularOpeningHours,places.attributions'
LOCK = threading.RLock()
JOBS = {}
TTL = 3600


def default_queries():
    return sorted({n.strip() for n in categories().values() if n and n.strip() and '#REF!' not in n})


def validate_geometry(value):
    try:
        if len(json.dumps(value)) > 2_000_000:
            raise ValueError('Boundary is too large; simplify it to under 2 MB.')
        geom = shape(value)
        if geom.geom_type not in ('Polygon', 'MultiPolygon') or geom.is_empty or not geom.is_valid:
            raise ValueError('Select a valid Polygon or MultiPolygon without self-intersections.')
        w, s, e, n = geom.bounds
        if not all(math.isfinite(v) for v in (w,s,e,n)) or not (-180 <= w < e <= 180 and -85 <= s < n <= 85):
            raise ValueError('Use a WGS84 polygon between latitudes -85 and 85.')
        if e-w > 180:
            raise ValueError('Split polygons crossing the date line before collection.')
        if (e-w)*111.32*math.cos(math.radians((s+n)/2))*(n-s)*111.32 > 100:
            raise ValueError('Select an area with a bounding box below 100 square kilometres; split larger areas.')
        return geom
    except (TypeError, KeyError, AttributeError, ShapelyError) as exc:
        raise ValueError('Invalid GeoJSON polygon.') from exc


def tiles(geom, metres):
    w,s,e,n = geom.bounds
    dy = metres / 111320
    dx = dy / math.cos(math.radians((s+n)/2))
    nx,ny = math.ceil((e-w)/dx),math.ceil((n-s)/dy)
    if nx*ny > 2500:
        raise ValueError('Too many search cells; use a larger grid size or smaller area.')
    return [(w+i*dx,s+j*dy,min(e,w+(i+1)*dx),min(n,s+(j+1)*dy))
            for i in range(nx) for j in range(ny)
            if geom.intersection(box(w+i*dx,s+j*dy,min(e,w+(i+1)*dx),min(n,s+(j+1)*dy))).area > 0]


def search(provider, bounds, query, language, token=None):
    cfg = get_settings()
    w,s,e,n = bounds
    if provider == 'google':
        if not cfg.google_places_api_key:
            raise ProviderError('GOOGLE_PLACES_API_KEY is not configured.')
        payload = {'textQuery':query,'pageSize':20,'locationRestriction':{'rectangle':{
            'low':{'latitude':s,'longitude':w},'high':{'latitude':n,'longitude':e}}}}
        if token: payload['pageToken'] = token
        if language: payload['languageCode'] = language
        result = request_json('https://places.googleapis.com/v1/places:searchText',payload,
                              {'X-Goog-Api-Key':cfg.google_places_api_key.get_secret_value(),'X-Goog-FieldMask':FIELDS})
        return result.get('places',[]),result.get('nextPageToken')
    result = grab_search(bounds,query,language,token,storage=True)
    return result.get('ResultItems',[]),result.get('NextToken')



DAYS = ('SUNDAY','MONDAY','TUESDAY','WEDNESDAY','THURSDAY','FRIDAY','SATURDAY')  # Google day 0 = Sunday
ISO3 = {'PH':'PHL','ID':'IDN','TH':'THA','MY':'MYS','SG':'SGP','VN':'VNM','KH':'KHM','MM':'MMR','LA':'LAO','BN':'BRN','TL':'TLS'}
MOBILE = {'63':('9',),'62':('8',),'66':('6','8','9')}  # country code -> mobile prefixes that follow it
STATUS = {'OPERATIONAL':'ACTIVE','CLOSED_TEMPORARILY':'TEMPORARILY_CLOSED','CLOSED_PERMANENTLY':'CLOSED'}


def label(value):
    return (value or '').replace('_',' ').strip().title()


def here_category(labels, query=None):
    """HERE (code, name, method): provider labels by HERE name, then HERE Find Places code; else the searched HERE category."""
    names, places = category_by_name(), find_places()
    for text in labels:
        key = norm_name(text or '')
        code = names.get(key) or places.get(key)
        if code: return code, categories().get(code,''), 'provider category'
    code = names.get(norm_name(query or ''))
    return (code, categories()[code], 'search term') if code else ('', '', '')


def food(labels):
    """HERE food type (code, name) from cuisine labels such as Google 'chinese_restaurant' or Grab 'chinese'."""
    lookup = foodtype_by_name()
    for text in labels:
        key = norm_name(re.sub(r'(?i)[ _]?(restaurant|food|cuisine)$','',text or ''))
        if key and key in lookup: return lookup[key], foodtypes()[lookup[key]]
    return '', ''


def chain(name, country):
    """Longest HERE chain the POI name equals or starts with; chainid '0' when none, as in the HERE extracts."""
    key, best, length = norm_name(name), ('0',''), 0
    for chain_name,cid in chain_ids(country).items():
        c = norm_name(chain_name)
        if c and (key == c or key.startswith(c+' ')) and len(c) > length: best, length = (cid,chain_name), len(c)
    if best[0] == '0':  # a HERE chain whose chainid is not in the extracts
        known = [c for c in chain_names(country) if norm_name(c) and (key == norm_name(c) or key.startswith(norm_name(c)+' '))]
        if known: best = ('',max(known,key=len))
    return best


def phones(numbers):
    """Digits with country code, as HERE stores them, split into (phone_numbers, mobile_numbers)."""
    fixed, mobile = [], []
    for n in numbers:
        digits = re.sub(r'\D','',n or '')
        if not digits: continue
        cc = next((c for c in MOBILE if digits.startswith(c)), None)
        (mobile if cc and digits[len(cc):].startswith(MOBILE[cc]) else fixed).append(digits)
    return ','.join(dict.fromkeys(fixed)), ','.join(dict.fromkeys(mobile))


def hours(opening):
    """Google regularOpeningHours.periods -> HERE 'MONDAY:08:30:00-18:00:00,...' (open days only, split at midnight)."""
    periods = (opening or {}).get('periods') or []
    if len(periods) == 1 and 'close' not in periods[0]:  # Google's 24/7 form
        return ','.join(f'{d}:00:00:00-23:59:59' for d in DAYS[1:]+DAYS[:1])
    slots = {d:[] for d in range(7)}
    for p in periods:
        o, c = p.get('open') or {}, p.get('close') or {}
        if 'day' not in o or 'day' not in c: continue
        start, end = (o.get('hour',0),o.get('minute',0)), (c.get('hour',0),c.get('minute',0))
        if c['day'] == o['day'] and end > start:
            slots[o['day']].append((start,end))
        else:  # runs past midnight
            slots[o['day']].append((start,None))
            if end != (0,0): slots[c['day']].append(((0,0),end))
    fmt = lambda t: '23:59:59' if t is None else f'{t[0]:02d}:{t[1]:02d}:00'
    return ','.join(f'{DAYS[d]}:{fmt(a)}-{fmt(b)}' for d in (1,2,3,4,5,6,0) for a,b in sorted(slots[d],key=lambda x: x[0]))


def normalized(provider, item, query=None):
    r = {h:'' for h in HEADERS}
    r['suppliers'] = '[GOOGLE_PLACES]' if provider == 'google' else '[GRABMAPS]'
    if provider == 'google':
        loc = item.get('location',{})
        lon,lat = loc.get('longitude'),loc.get('latitude')
        phone,mobile = phones([item.get('internationalPhoneNumber') or item.get('nationalPhoneNumber','')])
        r.update(placeid=item.get('id',''),name1=(item.get('displayName') or {}).get('text',''),
                 upaddress=item.get('formattedAddress',''),status=STATUS.get(item.get('businessStatus'),''),
                 phone_numbers=phone,mobile_numbers=mobile,webaddress=item.get('websiteUri',''),
                 opening_hours=hours(item.get('regularOpeningHours')))
        mapping = {'street_number':'hno','route':'streetname','postal_code':'postalcode',
                   'administrative_area_level_1':'admin2','administrative_area_level_2':'admin3','locality':'admin4',
                   'sublocality_level_1':'admin5','sublocality':'admin5','neighborhood':'admin5'}
        for c in item.get('addressComponents',[]):
            for kind in c.get('types',[]):
                if kind in mapping and not r[mapping[kind]]: r[mapping[kind]] = c.get('longText','')
                if kind == 'country': r['country'] = ISO3.get(c.get('shortText'),'')
        provider_label = (item.get('primaryTypeDisplayName') or {}).get('text') or label(item.get('primaryType'))
        labels = [provider_label,label(item.get('primaryType')),*map(label,item.get('types',[]))]
        cuisine = [label(t) for t in [item.get('primaryType'),*item.get('types',[])] if t]
        attribution = json.dumps(item.get('attributions',[]),ensure_ascii=False)
    else:
        p = item.get('Address') or {}
        lon,lat = (item.get('Position') or [None,None])[:2]
        cats = item.get('Categories') or []
        primary = next((c for c in cats if c.get('Primary')),cats[0] if cats else {})
        access = next((a.get('Position') for a in item.get('AccessPoints') or [] if a.get('Position')), None)
        r.update(placeid=item.get('PlaceId',''),name1=item.get('Title',''),upaddress=p.get('Label',''),
                 country=(p.get('Country') or {}).get('Code3',''),
                 hno=p.get('AddressNumber',''),streetname=p.get('Street',''),postalcode=p.get('PostalCode',''),
                 admin2=(p.get('Region') or {}).get('Name') or p.get('Locality',''),
                 admin3=(p.get('SubRegion') or {}).get('Name',''),
                 admin4=p.get('District',''),admin5=p.get('SubDistrict',''))
        if access: r.update(routinglongitude=access[0],routinglatitude=access[1])
        # Id is Grab's English category (casual_dining); Name is often the cuisine; LocalizedName is the local language.
        provider_label = label(primary.get('Id')) or primary.get('Name','')
        labels = [label(primary.get('Id')),primary.get('Name'),*(label(c.get('Id')) for c in cats)]
        cuisine = [primary.get('Name'),*(f.get('Name') or f.get('LocalizedName') for f in item.get('FoodTypes') or [])]
        attribution = 'GrabMaps via Amazon Location Service Places v2.'
    if not r['placeid'] or lon is None or lat is None:
        return None
    if not all(isinstance(v,(int,float)) and math.isfinite(v) for v in (lon,lat)):
        return None
    code,name,method = here_category(labels,query)
    r.update(primarycategory=code,primarycategorynm=name or provider_label)
    attribution += f' | HERE category from {method or "no match"}; provider category: {provider_label or "none"}.'
    if code.startswith(('100-','200-')):
        r['food_id'],r['food_name'] = food(cuisine)
    r['chainid'],r['chainname'] = chain(r['name1'],r['country'])
    r.update(displaylatitude=round(lat,6),displaylongitude=round(lon,6))
    return r,attribution


def public(job):
    with LOCK:
        return json.loads(json.dumps({k:v for k,v in job.items() if k not in ('owner','rows','cancel','geometry')}))


def start(owner, name, geometry, queries, budget, grid, language, providers=('google','grab')):
    geom = validate_geometry(geometry)
    cells = tiles(geom,grid)
    with LOCK:
        now = time.time()
        for key in list(JOBS):
            if JOBS[key]['state'] != 'running' and now-JOBS[key].get('finished_at',JOBS[key]['created']) > TTL:
                del JOBS[key]
        if any(j['state']=='running' for j in JOBS.values()):
            raise ValueError('Another collection is running. Wait for it to finish or cancel your run.')
        while len(JOBS) >= 20:
            del JOBS[next(iter(JOBS))]
        job = {'id':str(uuid4()),'owner':str(owner),'name':name,'created':now,'state':'running',
               'geometry':geometry,'queries':queries,'language':language,'grid_metres':grid,
               'cells':len(cells),'budget_per_provider':budget,'cancel':threading.Event(),'rows':{},
               'providers':{p:{'state':'queued','requests':0,'records':0,'duplicates':0,'outside':0,
                               'invalid_results':0,'capped_queries':0,'remaining_queries':len(cells)*len(queries),'error':None}
                            for p in providers},
               'notice':'Search results are not exhaustive. This is a provider extract, not an approved HERE delivery. Download within one hour; server restart removes runs.'}
        JOBS[job['id']] = job
    threading.Thread(target=run,args=(job,geom,cells,budget,language),daemon=True).start()
    return public(job)


def run(job, geom, cells, budget, language):
    try:
        for provider in job['providers']:
            stats = job['providers'][provider]
            found = {}
            # Sweep each category across all cells; recursively revisit capped cells.
            queue = deque((b,q,0,None,0) for q in job['queries'] for b in cells)
            with LOCK: stats['state'] = 'running'
            while queue and stats['requests'] < budget and not job['cancel'].is_set():
                bounds,query,depth,token,page = queue.popleft()
                with LOCK: stats['requests'] += 1
                try:
                    items,next_token = search(provider,bounds,query,language,token)
                    if not isinstance(items,list): raise ProviderError('Unexpected provider response.')
                except ProviderError as exc:
                    with LOCK: stats.update(error=str(exc),state='failed')
                    queue.appendleft((bounds,query,depth,token,page))
                    break
                for item in items:
                    record = normalized(provider,item,query)
                    if not record:
                        with LOCK: stats['invalid_results'] += 1
                        continue
                    row,attribution = record
                    if not geom.covers(Point(row['displaylongitude'],row['displaylatitude'])):
                        with LOCK: stats['outside'] += 1
                        continue
                    if row['placeid'] in found:
                        with LOCK: stats['duplicates'] += 1
                    else: found[row['placeid']] = (row,attribution)
                cap = (provider=='grab' and not next_token and len(items)>=100) or (provider=='google' and (page>=2 and (next_token or len(items)>=20)))
                if next_token and (provider=='grab' or page < 2):
                    queue.appendleft((bounds,query,depth,next_token,page+1))
                elif cap:
                    with LOCK: stats['capped_queries'] += 1
                    if depth < 2:
                        w,s,e,n = bounds; x,y = (w+e)/2,(s+n)/2
                        for child in ((w,s,x,y),(x,s,e,y),(w,y,x,n),(x,y,e,n)):
                            if geom.intersection(box(*child)).area > 0: queue.append((child,query,depth+1,None,0))
                with LOCK: stats.update(records=len(found),remaining_queries=len(queue))
                job['cancel'].wait(.15)
            with LOCK:
                job['rows'][provider] = list(found.values())
                stats.update(records=len(found),remaining_queries=len(queue))
                if stats['state'] != 'failed':
                    stats['state'] = 'cancelled' if job['cancel'].is_set() else 'limited' if queue else 'finished'
        with LOCK: job['state'] = 'cancelled' if job['cancel'].is_set() else 'finished'
    except Exception:
        with LOCK:
            job['state'] = 'failed'
            job['notice'] = 'Collection stopped unexpectedly. Available completed-provider results may be downloaded.'
    finally:
        with LOCK: job['finished_at'] = time.time()


def workbook(job):
    wb = Workbook(); wb.remove(wb.active)
    wb.properties.title = 'JSAN provider extract — incomplete coverage, not HERE delivery'
    # 'All POIs' combines both providers in one sheet (the suppliers column names the source).
    for providers,title in ((('google','grab'),'All POIs'),(('google',),'Google'),(('grab',),'Grab')):
        providers = [p for p in providers if p in job['providers']]
        ws = wb.create_sheet(title); ws.append(HEADERS)
        for cell in ws[1]:
            cell.font = Font(bold=True,color='FFFFFF'); cell.fill = PatternFill('solid',fgColor='005B96')
        summary = json.dumps({'providers':{p:job['providers'][p] for p in providers},'aoi':job['name'],'run_id':job['id'],
                              'retrieved_at':datetime.fromtimestamp(job['created'],timezone.utc).isoformat(),
                              'notice':job['notice'],
                              'mapping':'placeid is the provider ID, not a HERE ID. primarycategory, food_id and chainid come from the HERE reference files (method in each placeid comment). HERE-only fields (cqs, linkpvid, pointaddress, catalog, version) and fields the provider does not supply are blank.'},ensure_ascii=False)
        ws['A1'].comment = Comment(summary,'JSAN')
        for row,attribution in (r for p in providers for r in job['rows'].get(p,[])):
            # Strip control characters before append: openpyxl rejects them when the cell is created.
            ws.append([ILLEGAL_CHARACTERS_RE.sub('',v)[:32767] if isinstance(v,str) else v for v in (row.get(h,'') for h in HEADERS)])
            for cell in ws[ws.max_row]:
                if isinstance(cell.value,str):
                    cell.data_type = 's'  # Never execute provider strings as Excel formulae.
            ws.cell(ws.max_row,1).comment = Comment(ILLEGAL_CHARACTERS_RE.sub('',attribution),'Source attribution')
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for column in ws.columns: ws.column_dimensions[column[0].column_letter].width=22
    out=io.BytesIO(); wb.save(out); return out.getvalue()
