"""Strict, atomic imports. Original field names and values remain available in JSONB."""
import csv
import io
import math
import re
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook

AOIS = {
    "Divisoria": "PHL",
    "Kampung Bali": "IDN",
    "Thanon Phaya Thai": "THA",
}
MAX_BYTES = 25 * 1024 * 1024
MAX_ROWS = 100_000


class ImportInvalid(ValueError):
    def __init__(self, message, issues=None):
        super().__init__(message)
        self.issues = issues or []


def key(value):
    return re.sub(r"[^a-z0-9]", "", str(value).lower())


def string(value):
    if value is None:
        return ""
    return value.isoformat() if hasattr(value, "isoformat") else str(value).strip()


def parse_file(content: bytes, filename: str, aoi: str):
    if aoi not in AOIS:
        raise ImportInvalid("Select one of the three assigned AOIs.")
    if len(content) > MAX_BYTES:
        raise ImportInvalid("Maximum file size is 25 MB.")
    ext = filename.rsplit(".", 1)[-1].lower()
    tables = []
    workbook = None
    try:
        if ext == "csv":
            try:
                source = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                raise ImportInvalid("Save the CSV with UTF-8 encoding and upload it again.")
            reader = csv.reader(io.StringIO(source))
            tables.append(("CSV", next(reader, []), reader))
        elif ext == "xlsx":
            with ZipFile(io.BytesIO(content)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
                    raise ImportInvalid("The expanded workbook exceeds 100 MB.")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            for sheet in workbook:
                if sheet.sheet_state != "visible":
                    continue
                rows = sheet.iter_rows(values_only=True)
                tables.append((sheet.title, next(rows, []), rows))
        else:
            raise ImportInvalid("Supported formats are .csv and .xlsx. Save older Excel files as .xlsx.")

        records, columns, issues, seen = [], [], [], set()
        count = 0
        for sheet, header, rows in tables:
            headers = [string(v) for v in header]
            # Excel templates often have formatting-only columns beyond the table.
            while headers and not headers[-1]:
                headers.pop()
            if not headers:
                continue
            if any(not h for h in headers) or len(set(map(key, headers))) != len(headers):
                raise ImportInvalid(f"{sheet}: headers must be nonblank and unique.")
            normalized = {key(h) for h in headers}
            if not normalized.intersection({"placeid", "uniqueid", "poiid"}):
                raise ImportInvalid(f"{sheet}: missing placeid / unique_id column.")
            for h in headers:
                if h not in columns:
                    columns.append(h)
            for rownum, values in enumerate(rows, 2):
                if not any(string(v) for v in values):
                    continue
                count += 1
                if count > MAX_ROWS:
                    raise ImportInvalid("Maximum import size is 100,000 records.")
                try:
                    if len(values) > len(headers) and any(string(v) for v in values[len(headers):]):
                        raise ValueError("More values than column headers")
                    raw = {h: string(values[i]) if i < len(values) else "" for i, h in enumerate(headers)}
                    record = normalize(raw, AOIS[aoi])
                    if record["place_id"] in seen:
                        raise ValueError(f"Duplicate place ID: {record['place_id']}")
                    seen.add(record["place_id"])
                    records.append(record)
                except ValueError as exc:
                    if len(issues) < 50:
                        issues.append({"sheet": sheet, "row": rownum, "message": str(exc)})
        if issues:
            raise ImportInvalid("Import rejected. Fix the reported rows and retry; no records were saved. Showing up to 50 issues.", issues)
        if not records:
            raise ImportInvalid("The file has no populated POI rows.")
        return records, columns
    except ImportInvalid:
        raise
    except (BadZipFile, KeyError, ValueError, TypeError, csv.Error) as exc:
        raise ImportInvalid("Cannot read this file. Check the CSV/Excel format and headers.") from exc
    finally:
        if workbook:
            workbook.close()


def normalize(raw, country):
    row = {key(k): v for k, v in raw.items()}

    def pick(*names):
        return next((row[n] for n in names if row.get(n)), None)

    place_id = pick("placeid", "uniqueid", "poiid")
    name = pick("name1", "poiname", "name")
    if not place_id or not name:
        raise ValueError("Place ID and POI name are required")
    provided_country = pick("country")
    if provided_country and provided_country.upper() != country:
        raise ValueError(f"Country must match the selected AOI ({country})")

    def point(lat, lon, optional=False):
        if optional and lat is None and lon is None:
            return None, None
        try:
            lat, lon = float(lat), float(lon)
        except (TypeError, ValueError):
            raise ValueError("A complete numeric latitude/longitude pair is required")
        if not math.isfinite(lat) or not math.isfinite(lon) or abs(lat) > 90 or abs(lon) > 180:
            raise ValueError("Coordinates must be valid WGS84 decimal degrees")
        return lat, lon

    lat, lon = point(pick("displaylatitude", "displaypointlatitude", "latitude", "lat"),
                     pick("displaylongitude", "displaypointlongitude", "longitude", "lon", "lng"))
    rlat, rlon = point(pick("routinglatitude", "routinglat"), pick("routinglongitude", "routinglon"), True)
    return {
        "place_id": place_id, "name": name,
        "category_code": pick("primarycategory", "categorycode", "category"),
        "category_name": pick("primarycategorynm", "categoryname", "category"),
        "country": country, "address": pick("upaddress", "unparsedaddressname", "address"),
        "latitude": lat, "longitude": lon, "routing_latitude": rlat, "routing_longitude": rlon,
        "raw": raw,
    }
