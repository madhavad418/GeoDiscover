"""Build compact runtime reference data from the HERE POC package.

Usage (from backend/, dev dependencies installed):
    python build_reference_data.py "path/to/POC Material"

Writes app/resources/reference_data.json and copies the three delivery templates to
app/resources/templates/. Re-run whenever HERE issues new reference files or templates.
"""
import argparse
import csv
import hashlib
import json
import shutil
from pathlib import Path

from openpyxl import load_workbook

RESOURCES = Path(__file__).resolve().parent / "app" / "resources"
COUNTRIES = {"PHL", "IDN", "THA"}
TEMPLATES = {
    "new": "Supplier Name_POI_New_DDMMYY.xlsx",
    "update": "Supplier Name_POI_Updated_DDMMYY.xlsx",
    "delete": "Supplier Name_POI_Deleted_DDMMYY.xlsx",
}


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def rows(path, sheet=None):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it)]
    for values in it:
        if any(v not in (None, "") for v in values):
            yield dict(zip(header, values))
    wb.close()


def categories(path):
    names = {}
    for sheet in ("PDS-to-POIXML_CatSys", "PDS-to-OPR_CatSys", "POIXML-to-PDS_CatSys", "OPR-to-PDS_CatSys"):
        for row in rows(path, sheet):
            code, name = row.get("PDS Category Code"), row.get("PDS Category Name")
            if code and isinstance(code, str) and len(code) == 13 and code[3] == "-":
                names.setdefault(code.strip(), str(name or "").strip())
    return dict(sorted(names.items()))


def find_places(path):
    """OPR Find Places codes (e.g. casual-dining, hostel) -> PDS category code: a crosswalk for provider category IDs."""
    out = {}
    for row in rows(path, "OPR-to-PDS_CatSys"):
        code, pds = str(row.get("Find Places Code") or "").strip(), str(row.get("PDS Category Code") or "").strip()
        if code and len(pds) == 13:
            out.setdefault(code.replace("-", " "), pds)
    for row in rows(path, "PDS-to-OPR_CatSys"):  # fill Find Places codes missing above with their first PDS code
        code, pds = str(row.get("Find Places Code") or "").strip(), str(row.get("PDS Category Code") or "").strip()
        if code and len(pds) == 13:
            out.setdefault(code.replace("-", " "), pds)
    return dict(sorted(out.items()))


def foodtypes(path):
    out = {}
    for row in rows(path, "PDS-to-POIXML_Foodtype"):
        code, name = str(row.get("PDS Foodtype Code") or "").strip(), str(row.get("PDS Foodtype") or "").strip()
        if code and name:
            out[code] = name
    return out


def chain_ids(extracts):
    """Chain name -> HERE chainid per country, as used in the supplied HERE baseline extracts."""
    out = {c: {} for c in COUNTRIES}
    for path in sorted(extracts.glob("*.csv")):
        with path.open(encoding="utf-8-sig", newline="") as f:
            for row in csv.DictReader(f):
                country, cid, name = row.get("country", ""), (row.get("chainid") or "").strip(), (row.get("chainname") or "").strip()
                if country in out and name and cid not in ("", "0"):
                    out[country].setdefault(name, cid)
    return {c: dict(sorted(v.items())) for c, v in out.items()}


def profanity(path):
    out = []
    for row in rows(path):
        phrase = str(row.get("prohibitedPhrase") or "").strip()
        scope = {s.strip() for s in str(row.get("countryLanguage") or "").split(",") if s.strip()}
        countries = sorted({s.split("-")[0] for s in scope} & (COUNTRIES | {"ALL"}))
        if phrase and countries:
            out.append({"phrase": phrase, "whole_name": bool(row.get("hasToBeEntirePlaceName")), "countries": countries})
    return out


def chains(path):
    out = {c: set() for c in COUNTRIES}
    for row in rows(path):
        country, name = str(row.get("COUNTRY") or ""), str(row.get("NAME") or "").strip()
        if country in COUNTRIES and name:
            out[country].add(name)
    return {c: sorted(v) for c, v in out.items()}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("package", type=Path)
    package = parser.parse_args().package
    sources = {
        "categories": package / "HERE Category Metadata-2.xlsx",
        "profanity": package / "Profanity_APR_2026.xlsx",
        "chains": package / "Chains MetaData-2.xlsx",
    }
    data = {
        "sources": {k: {"file": p.name, "sha256": sha256(p)} for k, p in sources.items()},
        "categories": categories(sources["categories"]),
        "profanity": profanity(sources["profanity"]),
        "chains": chains(sources["chains"]),
        "find_places": find_places(sources["categories"]),
        "foodtypes": foodtypes(sources["categories"]),
        "chain_ids": chain_ids(package / "Requested_AOI_Places" / "data_extract"),
        "templates": {},
    }
    (RESOURCES / "templates").mkdir(exist_ok=True)
    for action, name in TEMPLATES.items():
        src = package / name
        shutil.copyfile(src, RESOURCES / "templates" / name)
        data["templates"][action] = {"file": name, "sha256": sha256(src)}
    (RESOURCES / "reference_data.json").write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"categories={len(data['categories'])} profanity={len(data['profanity'])} "
          f"chains={ {k: len(v) for k, v in data['chains'].items()} } find_places={len(data['find_places'])} "
          f"foodtypes={len(data['foodtypes'])} chain_ids={ {k: len(v) for k, v in data['chain_ids'].items()} } templates={len(data['templates'])}")


if __name__ == "__main__":
    main()
